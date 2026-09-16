from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from io import BytesIO
import json
import sqlite3

import pytest
from openpyxl import load_workbook

from point_activite.storage import Store, filter_records, ConflictError
from point_activite.excel import export_excel


def entry(day='2026-09-15', **kwargs):
    return dict(date=day, theme='Formalités et réglementation', sub='',
                title='Procédure de création', content='Créer une activité avec Salesforce.',
                type='Information', action='', who='Équipe accueil', links='', note='',
                **kwargs)


def test_search_and_absence_boundaries():
    rows=[dict(entry(day),id=str(i)) for i,day in enumerate(
        ['2026-09-13','2026-09-14','2026-09-15','2026-09-16',None])]
    assert [r['id'] for r in filter_records(rows,start='2026-09-14',end='2026-09-16',end_inclusive=False)]==['2','1']
    assert len(filter_records(rows,start='2026-09-14',end='2026-09-16'))==3
    assert len(filter_records(rows,query='CREER salesforce',person='equipe'))==5
    assert filter_records(rows,query='Salesforce introuvable')==[]
    assert len(filter_records(rows,only_undated=True))==1
    assert filter_records(rows,themes=['Taxi et VTC'])==[]
    assert filter_records(rows,types=['Contact'])==[]
    assert filter_records(rows,start='2026-09-15',end='2026-09-15',end_inclusive=False)==[]


def test_seed_once_no_overwrite(tmp_path):
    (tmp_path/'seed').mkdir()
    (tmp_path/'seed/history.json').write_text(json.dumps({'records':[dict(entry(),id='PA-0001')]}))
    s=Store(tmp_path)
    r=s.all()[0]
    s.save(dict(r,title='Titre corrigé'),'Romain',r['id'],r['version'])
    other=Store(tmp_path)
    assert len(other.all())==1
    assert other.all()[0]['title']=='Titre corrigé'


def test_concurrent_edits_history_and_reads(tmp_path):
    s=Store(tmp_path)
    ident=s.save(entry(),'Alice')
    old=s.all()[0]
    s.mark_read('Alice',old)
    assert s.read_versions('Alice')[ident]==1
    def edit(author):
        try:
            return s.save(dict(old,title=author),author,ident,old['version'])
        except ConflictError:
            return 'conflict'
    with ThreadPoolExecutor(max_workers=2) as pool:
        results=list(pool.map(edit,['Bob','Claire']))
    assert results.count('conflict')==1
    new=s.all()[0]
    assert new['version']==2
    assert new['author']=='Alice'
    assert s.read_versions('Alice')[ident] != new['version']
    assert json.loads(s.revisions(ident)[0]['payload'])['title']==old['title']
    assert not s.read_versions('Bob')


def test_backup_contains_committed_records(tmp_path):
    s=Store(tmp_path/'live')
    s.save(entry(),'Alice')
    dest=tmp_path/'backup.sqlite3'
    dest.write_bytes(s.backup())
    with sqlite3.connect(dest) as db:
        assert db.execute('SELECT count(*) FROM entries').fetchone()[0]==1


def test_excel_dates_verbatim_text_actions_and_empty_export():
    rows=[dict(entry(), id='X',content='=HYPERLINK("bad")\nTexte & <exact>',
               action='Rappeler',status='À faire',attachments=[{'name':'note.pdf'}]),
          dict(entry(None),id='Y')]
    wb=load_workbook(BytesIO(export_excel(rows,'Test')))
    assert wb.sheetnames==['Base','Actions','Guide']
    base=wb['Base']
    assert base.max_row==3
    by_id={r[0].value:r for r in list(base.rows)[1:]}
    assert by_id['X'][1].value==datetime(2026,9,15)
    assert by_id['Y'][1].value is None
    assert by_id['X'][5].value==rows[0]['content']
    assert by_id['X'][5].data_type=='s'
    assert wb['Actions'].max_row==2
    assert base.freeze_panes=='A2'
    assert base.auto_filter.ref=='A1:Q3'
    assert load_workbook(BytesIO(export_excel([])))['Base'].max_row==1


def test_new_contribution_requires_identity_and_date(tmp_path):
    s=Store(tmp_path)
    with pytest.raises(ValueError): s.save(entry(),'')
    with pytest.raises(ValueError): s.save(entry(None),'Alice')
