from pathlib import Path
from datetime import date

from streamlit.testing.v1 import AppTest

from point_activite.storage import Store

APP=Path(__file__).resolve().parents[1]/'streamlit_app.py'


def app(tmp_path,monkeypatch):
    monkeypatch.setenv('POINT_ACTIVITE_DATA_DIR',str(tmp_path))
    monkeypatch.setenv('POINT_ACTIVITE_AUTH','local')
    Store(tmp_path).save(dict(date='2026-09-15',title='Formation Salesforce',
        content='Un échange pour l’équipe accueil.',theme='Outils et traitement des demandes',
        type='Action à réaliser',action='Consulter la procédure',status='À faire',who='Accueil'),'Alice')
    return AppTest.from_file(str(APP),default_timeout=20).run()


def test_navigation_search_and_absence(tmp_path,monkeypatch):
    at=app(tmp_path,monkeypatch)
    assert not at.exception
    for page in ['Base de connaissances','Retour d’absence','Actions','Ajouter une contribution','Exports et sauvegarde']:
        at.radio(key='page').set_value(page).run()
        assert not at.exception, page
    at.radio(key='page').set_value('Base de connaissances').run()
    at.text_input(key='search_query').set_value('salesforce').run()
    assert len(at.expander)==1
    at.text_input(key='search_query').set_value('inexistant').run()
    assert len(at.expander)==0


def test_contribute_edit_and_read(tmp_path,monkeypatch):
    at=app(tmp_path,monkeypatch)
    at.text_input(key='actor').set_value('Romain').run()
    at.radio(key='page').set_value('Ajouter une contribution').run()
    at.text_input(key='new_title').set_value('Nouvelle information')
    at.text_area(key='new_content').set_value('Une consigne complète.')
    submit=next(b for b in at.button if b.label=='Enregistrer la contribution')
    submit.click().run()
    assert not at.exception
    s=Store(tmp_path)
    assert len(s.all())==2
    at.radio(key='page').set_value('Base de connaissances').run()
    r=next(r for r in s.all() if r['title']=='Nouvelle information')
    at.button(key='search_read_'+r['id']).click().run()
    assert s.read_versions('Romain')[r['id']]==1
    at.button(key='search_edit_'+r['id']).click().run()
    assert not at.exception
    at.text_input(key=r['id']+'_title').set_value('Information corrigée')
    next(b for b in at.button if b.label=='Enregistrer la contribution').click().run()
    assert not at.exception
    assert next(e for e in s.all() if e['id']==r['id'])['title']=='Information corrigée'


def test_invalid_auth_mode_loads_no_data(tmp_path,monkeypatch):
    monkeypatch.setenv('POINT_ACTIVITE_DATA_DIR',str(tmp_path/'uncreated'))
    monkeypatch.setenv('POINT_ACTIVITE_AUTH','mistake')
    at=AppTest.from_file(str(APP)).run()
    assert at.error
    assert not (tmp_path/'uncreated/point_activite.sqlite3').exists()


def test_search_dropdown_contains_all_results_and_opens_last(tmp_path,monkeypatch):
    at=app(tmp_path,monkeypatch)
    store=Store(tmp_path)
    for i in range(25):
        store.save(dict(date='2026-09-14',title=f'Fiche test {i:02}',content=f'Contenu complet {i:02}',
                        theme='À qualifier',type='Information'),'Alice')
    at.radio(key='page').set_value('Base de connaissances').run()
    selector=at.selectbox(key='search_record')
    assert len(selector.options)==26
    assert len(at.number_input)==0
    last=sorted(store.all(),key=lambda r:(r['date'],r['created_at'],r['id']),reverse=True)[-1]
    selector.set_value(last['id']).run()
    assert not at.exception
    assert last['content'] in [t.value for t in at.text]
    at.text_input(key='search_query').set_value('Salesforce').run()
    assert len(at.selectbox(key='search_record').options)==1
    assert not at.exception
    at.text_input(key='actor').set_value('Alice').run()
    chosen=next(r for r in store.all() if r['title']=='Formation Salesforce')
    at.button(key='search_read_'+chosen['id']).click().run()
    at.checkbox(key='search_unread').set_value(True).run()
    assert not at.exception
    assert not [s for s in at.selectbox if s.key=='search_record']


def test_absence_calendar_includes_last_day_and_handles_partial_range(tmp_path,monkeypatch):
    at=app(tmp_path,monkeypatch)
    at.radio(key='page').set_value('Retour d’absence').run()
    at.date_input(key='absence_period').set_value((date(2026,9,15),date(2026,9,15))).run()
    assert not at.exception
    assert len(at.expander)==1  # A single-day holiday includes that day's contribution.
    assert any('1 contribution(s)' in m.value for m in at.success)
    at.date_input(key='absence_period').set_value((date(2026,9,15),)).run()
    assert not at.exception
    assert len(at.expander)==0
    assert any('date de fin' in m.value for m in at.info)
    at.button(key='absence_period_La semaine dernière').click().run()
    assert not at.exception
    start,end=at.date_input(key='absence_period').value
    assert start.weekday()==0 and end.weekday()==6 and (end-start).days==6


def test_period_export_includes_boundaries_and_total_keeps_undated(tmp_path,monkeypatch):
    from point_activite import excel
    from io import BytesIO
    from openpyxl import load_workbook
    exported=[]
    original=excel.export_excel
    def capture(records,description):
        result=original(records,description)
        exported.append((description,load_workbook(BytesIO(result))))
        return result
    monkeypatch.setattr(excel,'export_excel',capture)
    at=app(tmp_path,monkeypatch)
    s=Store(tmp_path)
    for day in ['2026-09-13','2026-09-14','2026-09-16','2026-09-17']:
        s.save(dict(date=day,title=day,content='Test période',theme='À qualifier',type='Information'),'Alice')
    ident=s.save(dict(date='2026-09-18',title='Sans date',content='Archive',theme='À qualifier',type='Information'),'Alice')
    r=next(r for r in s.all() if r['id']==ident)
    s.save(dict(r,date=None),'Alice',ident,r['version'])
    at.radio(key='page').set_value('Exports et sauvegarde').run()
    assert exported[-1][1]['Base'].max_row==7
    at.radio(key='export_scope').set_value('De date à date').run()
    at.date_input(key='export_period').set_value((date(2026,9,14),date(2026,9,16))).run()
    assert not at.exception
    description,wb=exported[-1]
    assert wb['Base'].max_row==4
    assert {r[1].value.date() for r in list(wb['Base'].rows)[1:]}=={date(2026,9,14),date(2026,9,15),date(2026,9,16)}
    assert '2026-09-14' in description and '2026-09-16' in description
    count=len(exported)
    at.date_input(key='export_period').set_value((date(2026,9,14),)).run()
    assert not at.exception
    assert len(exported)==count  # No misleading download for an incomplete period.
    at.radio(key='export_scope').set_value('Toute la base').run()
    assert exported[-1][1]['Base'].max_row==7
